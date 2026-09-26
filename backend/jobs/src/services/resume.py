"""Resume upload handling: type checks, size limit, text extraction and atomic storage.

Only one resume is kept, as ``<resume_folder>/resume.pdf`` or ``resume.docx``. The client's
filename is used for nothing but the extension check.
"""

import io
import logging
import os
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import BinaryIO

logger = logging.getLogger(__name__)

MAX_RESUME_BYTES = 10 * 1024 * 1024
READ_CHUNK_BYTES = 1024 * 1024
RESUME_BASENAME = "resume"
RESUME_TYPES: tuple[str, ...] = ("pdf", "docx")

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"
DOCX_MAIN_PART = "word/document.xml"


class ResumeError(Exception):
    """A resume upload that cannot be accepted; ``status_code`` maps it to an HTTP response."""

    status_code = 400

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


class ResumeTooLargeError(ResumeError):
    status_code = 413


class UnsupportedResumeTypeError(ResumeError):
    status_code = 415


class UnreadableResumeError(ResumeError):
    status_code = 422


@dataclass(frozen=True)
class StoredResume:
    path: Path
    file_type: str
    uploaded_at: datetime


def read_limited(stream: BinaryIO, limit: int = MAX_RESUME_BYTES) -> bytes:
    """Read the upload in chunks, stopping as soon as it exceeds ``limit`` bytes."""
    buffer = bytearray()
    while chunk := stream.read(READ_CHUNK_BYTES):
        buffer.extend(chunk)
        if len(buffer) > limit:
            raise ResumeTooLargeError(
                f"The resume is larger than {limit // (1024 * 1024)} MB. Upload a smaller file."
            )
    return bytes(buffer)


def detect_file_type(filename: str | None, content: bytes) -> str:
    """``pdf`` or ``docx`` when both the extension and the file content agree."""
    extension = Path(filename or "").suffix.lower().lstrip(".")
    if extension == "pdf" and content.startswith(PDF_MAGIC):
        return "pdf"
    if extension == "docx" and content.startswith(ZIP_MAGIC) and _has_docx_body(content):
        return "docx"
    raise UnsupportedResumeTypeError("Upload the resume as a PDF (.pdf) or Word (.docx) file.")


def _has_docx_body(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return DOCX_MAIN_PART in archive.namelist()
    except zipfile.BadZipFile:
        return False


def extract_text(file_type: str, content: bytes) -> str:
    """Plain text of the resume; raises ``UnreadableResumeError`` when there is none."""
    try:
        text = _extract_pdf(content) if file_type == "pdf" else _extract_docx(content)
    except Exception as error:
        logger.warning("Resume text extraction failed type=%s error=%s", file_type, error)
        raise UnreadableResumeError(
            "The resume could not be read. Check that the file is not damaged or password "
            "protected."
        ) from None
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise UnreadableResumeError(
            "No text could be found in the resume. Scanned documents are not supported; upload "
            "a PDF or Word file that contains selectable text."
        )
    return text


def _extract_pdf(content: bytes) -> str:
    from pdfminer.high_level import extract_text as pdf_text

    return pdf_text(io.BytesIO(content))


def _extract_docx(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def resume_path(folder: Path, file_type: str) -> Path:
    return folder / f"{RESUME_BASENAME}.{file_type}"


def store_resume(folder: Path, file_type: str, content: bytes) -> StoredResume:
    """Atomically write ``resume.<type>`` and remove the resume with the other extension."""
    folder.mkdir(parents=True, exist_ok=True)
    target = resume_path(folder, file_type)
    descriptor, temp_name = tempfile.mkstemp(dir=folder, prefix=".resume-", suffix=".tmp")
    try:
        with os.fdopen(descriptor, "wb") as temp_file:
            temp_file.write(content)
            temp_file.flush()
            os.fsync(temp_file.fileno())
        os.replace(temp_name, target)
    except BaseException:
        Path(temp_name).unlink(missing_ok=True)
        raise
    for other_type in RESUME_TYPES:
        if other_type != file_type:
            resume_path(folder, other_type).unlink(missing_ok=True)
    return StoredResume(path=target, file_type=file_type, uploaded_at=_modified_at(target))


def delete_resume(folder: Path) -> None:
    for file_type in RESUME_TYPES:
        resume_path(folder, file_type).unlink(missing_ok=True)


def stored_resume(location: str | None) -> StoredResume | None:
    """Metadata for the resume at ``location`` if the file still exists."""
    if not location:
        return None
    path = Path(location)
    file_type = path.suffix.lower().lstrip(".")
    if file_type not in RESUME_TYPES or not path.is_file():
        return None
    return StoredResume(path=path, file_type=file_type, uploaded_at=_modified_at(path))


def _modified_at(path: Path) -> datetime:
    return datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
